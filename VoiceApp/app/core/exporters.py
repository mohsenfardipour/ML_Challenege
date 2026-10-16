"""Save transcripts as TXT (UTF-8), DOCX (right-to-left) or SRT subtitles."""

from __future__ import annotations

from pathlib import Path

from app.core.transcriber import Segment

DOCX_FONT = "Tahoma"   # installed on every Windows PC and has Persian glyphs (Vazirmatn is used if you set it)


def export_txt(path: Path, text: str) -> None:
    # utf-8-sig: the BOM makes Windows Notepad and Excel detect UTF-8 reliably.
    Path(path).write_text(text.strip() + "\n", encoding="utf-8-sig")


def _srt_time(seconds: float) -> str:
    ms = int(round(max(0.0, seconds) * 1000))
    h, ms = divmod(ms, 3_600_000)
    m, ms = divmod(ms, 60_000)
    s, ms = divmod(ms, 1000)
    return f"{h:02d}:{m:02d}:{s:02d},{ms:03d}"


def export_srt(path: Path, segments: list[Segment]) -> None:
    blocks = []
    for i, seg in enumerate((s for s in segments if s.text.strip()), 1):
        # U+202B (right-to-left embedding) helps players show Persian lines correctly.
        blocks.append(f"{i}\n{_srt_time(seg.start)} --> {_srt_time(seg.end)}\n\u202b{seg.text.strip()}\u202c\n")
    Path(path).write_text("\n".join(blocks), encoding="utf-8-sig")


# Word requires child elements in schema order; these are the ones that come after.
_PPR_AFTER_BIDI = ("w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing",
                   "w:mirrorIndents", "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment",
                   "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr",
                   "w:pPrChange")
_RPR_AFTER_SZCS = ("w:highlight", "w:u", "w:effect", "w:bdr", "w:shd", "w:fitText", "w:vertAlign",
                   "w:rtl", "w:cs", "w:em", "w:lang", "w:eastAsianLayout", "w:specVanish", "w:oMath")


def export_docx(path: Path, text: str, title: str | None = None, font: str = DOCX_FONT) -> None:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    doc = Document()

    def make_rtl(paragraph, size_pt: int, bold: bool = False):
        # Right-to-left paragraph; a bidi paragraph is right-aligned by default.
        ppr = paragraph._p.get_or_add_pPr()
        bidi = OxmlElement("w:bidi")
        ppr.insert_element_before(bidi, *_PPR_AFTER_BIDI)
        for run in paragraph.runs:
            run.bold = bold
            run.font.size = Pt(size_pt)
            run.font.name = font
            run.font.rtl = True
            rpr = run._r.get_or_add_rPr()
            fonts = rpr.get_or_add_rFonts()
            for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
                fonts.set(qn(attr), font)
            szcs = OxmlElement("w:szCs")
            szcs.set(qn("w:val"), str(size_pt * 2))
            rpr.insert_element_before(szcs, *_RPR_AFTER_SZCS)
            lang = OxmlElement("w:lang")
            lang.set(qn("w:bidi"), "fa-IR")
            rpr.insert_element_before(lang, "w:eastAsianLayout", "w:specVanish", "w:oMath")

    if title:
        make_rtl(doc.add_paragraph(title), 14, bold=True)
    for line in text.strip().splitlines():
        make_rtl(doc.add_paragraph(line), 12)
    doc.save(str(path))
